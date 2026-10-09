"""Module: @role: マクロ実行エンジンにおけるExcel COM/openpyxl操作、レコード抽出、ステータス更新を担う。"""

import ctypes
import logging
from pathlib import Path
import platform
import re
import time
from typing import Any, Callable, Dict, List, Optional, Set, Tuple

logger = logging.getLogger(__name__)


def col_idx_to_letter(col_idx: int) -> str:
    # Why: openpyxl依存なしで26列超(AA, AB等)のExcel列記号を確実に算出
    res = ""
    while col_idx > 0:
        col_idx, rem = divmod(col_idx - 1, 26)
        res = chr(65 + rem) + res
    return res


def extract_referenced_columns(commands_slice: List[Dict[str, Any]], header_map: Optional[Dict[str, str]] = None) -> Set[str]:
    # Why: 日本語ヘッダー名と列文字の双方を網羅検出し元データ列の上書き破壊を完全遮断
    ref_cols = set()
    pattern = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}|\$\{([^{}]+?)\}")

    def _scan(val: Any) -> None:
        if isinstance(val, str):
            for m in pattern.finditer(val):
                token = (m.group(1) or m.group(2)).strip()
                if "." in token:
                    token = token.split(".")[-1].strip()
                ref_cols.add(token.upper())
                if header_map:
                    for col_letter, h_name in header_map.items():
                        if token.lower() == str(h_name).lower():
                            ref_cols.add(col_letter.upper())
        elif isinstance(val, dict):
            for v in val.values():
                _scan(v)
        elif isinstance(val, list):
            for item in val:
                _scan(item)

    for c in commands_slice:
        _scan(c.get("args", {}))
    return ref_cols


def is_file_locked(path_str: str) -> bool:
    # Why: ファイル排他ロックの事前検知によりopenpyxl PermissionErrorを完全回避
    if not Path(path_str).exists():
        return False
    try:
        with open(path_str, "a+b"):
            pass
        return False
    except (PermissionError, IOError):
        return True


def read_excel_records(
    file_path: str,
    sheet_name: Optional[str] = None,
    start_row: int = 2,
    end_row: Optional[int] = None,
    status_col: Optional[str] = None,
    skip_completed: bool = True,
    target_dir: Optional[Path] = None,
) -> List[Dict[str, Any]]:
    # Why: 相対パス指定時もtarget_dirからのフォールバック解決によりファイル未検出を根絶
    if not file_path:
        return []
    p = Path(file_path)
    if not p.is_absolute() and target_dir and (target_dir / file_path).exists():
        resolved = str((target_dir / file_path).resolve())
    else:
        resolved = str(p.resolve())
    if not Path(resolved).exists():
        logger.error(f"Excel file not found: {resolved}")
        return []

    records = []
    try:
        import openpyxl
        wb = openpyxl.load_workbook(resolved, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        header_map = {}
        for col_idx in range(1, ws.max_column + 1):
            col_letter = col_idx_to_letter(col_idx)
            val = ws.cell(row=1, column=col_idx).value
            if val is not None:
                header_map[col_letter] = str(val).strip()

        max_r = end_row or ws.max_row
        for r in range(start_row, max_r + 1):
            row_has_val = False
            row_dict = {"_row_idx": r}
            for col_idx in range(1, ws.max_column + 1):
                col_letter = col_idx_to_letter(col_idx)
                raw_v = ws.cell(row=r, column=col_idx).value
                if raw_v is not None and str(raw_v).strip() != "":
                    row_has_val = True
                # Why: Excel数値のfloat化(100 -> 100.0)や指数表記を防ぎコード値の整合性を担保
                if raw_v is None:
                    v_str = ""
                elif hasattr(raw_v, "strftime"):
                    v_str = raw_v.strftime("%Y-%m-%d")
                elif isinstance(raw_v, float) and raw_v.is_integer():
                    v_str = str(int(raw_v))
                else:
                    v_str = str(raw_v).strip()

                row_dict[col_letter] = v_str
                row_dict[f"列{col_letter}"] = v_str
                row_dict[f"{col_letter}列"] = v_str
                if col_letter in header_map:
                    row_dict[header_map[col_letter]] = v_str

            if not row_has_val:
                continue
            if skip_completed and status_col:
                cur_status = str(row_dict.get(status_col.upper(), "")).strip()
                # Why: 中断・再テスト時のスタックを防ぐため「処理中」は未完了扱いで再実行
                if cur_status in ["完了", "DONE", "済", "スキップ", "SUCCESS"]:
                    continue
            records.append(row_dict)
        wb.close()
        return records
    except Exception as e:
        logger.warning(f"openpyxl headless read failed ({e}), falling back to COM...")

    if platform.system() == "Windows":
        excel = None
        should_quit = False
        try:
            import win32com.client
            # Why: 既に開かれているExcelからのインメモリ直接読み取りを優先し切断・排他ロック競合を完全回避
            try:
                excel = win32com.client.GetActiveObject("Excel.Application")
            except Exception:
                excel = win32com.client.Dispatch("Excel.Application")
                should_quit = True

            wb = None
            for open_wb in excel.Workbooks:
                if open_wb.FullName.lower() == resolved.lower() or open_wb.Name.lower() == Path(file_path).name.lower():
                    wb = open_wb
                    break

            should_close = False
            if wb is None:
                wb = excel.Workbooks.Open(resolved, ReadOnly=True)
                should_close = True

            ws = wb.Sheets(sheet_name) if sheet_name else wb.ActiveSheet
            vals = None
            for _ in range(3):
                try:
                    vals = ws.UsedRange.Value
                    break
                except Exception:
                    time.sleep(0.2)

            if should_close:
                wb.Close(SaveChanges=False)

            if vals and len(vals) >= start_row:
                headers = vals[0]
                header_map = {col_idx_to_letter(idx + 1): str(h).strip() for idx, h in enumerate(headers) if h is not None}
                max_r = min(end_row, len(vals)) if end_row else len(vals)
                for r in range(start_row, max_r + 1):
                    row_data = vals[r - 1]
                    if not any(c is not None and str(c).strip() != "" for c in row_data):
                        continue
                    row_dict = {"_row_idx": r}
                    for idx, cell in enumerate(row_data):
                        c_letter = col_idx_to_letter(idx + 1)
                        if cell is None:
                            c_str = ""
                        elif hasattr(cell, "strftime"):
                            c_str = cell.strftime("%Y-%m-%d")
                        elif isinstance(cell, float) and cell.is_integer():
                            c_str = str(int(cell))
                        else:
                            c_str = str(cell).strip()

                        row_dict[c_letter] = c_str
                        row_dict[f"列{c_letter}"] = c_str
                        row_dict[f"{c_letter}列"] = c_str
                        if c_letter in header_map:
                            row_dict[header_map[c_letter]] = c_str
                    if skip_completed and status_col:
                        cur_status = str(row_dict.get(status_col.upper(), "")).strip()
                        if cur_status in ["完了", "DONE", "済", "スキップ", "SUCCESS"]:
                            continue
                    records.append(row_dict)
        except Exception as com_err:
            logger.error(f"COM fallback read also failed: {com_err}")
        finally:
            if excel and should_quit:
                try:
                    excel.Quit()
                except Exception:
                    pass
                del excel
    return records


def write_excel_status(file_path: str, sheet_name: Optional[str], row_idx: int, status_col: str, status_text: str) -> None:
    # Why: 開かれているExcelはCOMインメモリ即時更新、非オープン時はopenpyxl直接保存を自動選択
    if not file_path or not status_col or not row_idx:
        return
    resolved = str(Path(file_path).resolve())
    locked = is_file_locked(resolved)

    if not locked:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(resolved)
            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            col_idx = openpyxl.utils.column_index_from_string(status_col)
            ws.cell(row=row_idx, column=col_idx, value=status_text)
            wb.save(resolved)
            wb.close()
            return
        except Exception as e:
            logger.warning(f"openpyxl status update failed ({e}), routing to COM...")

    if platform.system() == "Windows":
        excel = None
        should_quit = False
        try:
            import win32com.client
            try:
                excel = win32com.client.GetActiveObject("Excel.Application")
            except Exception:
                excel = win32com.client.Dispatch("Excel.Application")
                should_quit = True
            wb = None
            for open_wb in excel.Workbooks:
                if open_wb.FullName.lower() == resolved.lower() or open_wb.Name.lower() == Path(file_path).name.lower():
                    wb = open_wb
                    break
            should_close = False
            if wb is None:
                wb = excel.Workbooks.Open(resolved)
                should_close = True
            ws = wb.Sheets(sheet_name) if sheet_name else wb.ActiveSheet
            # Why: COMビジー時のRPC拒否に備え最大3回のリトライバックオフを実施
            for attempt in range(3):
                try:
                    ws.Range(f"{status_col}{row_idx}").Value = status_text
                    wb.Save()
                    break
                except Exception:
                    time.sleep(0.3)
            if should_close:
                wb.Close(SaveChanges=True)
        except Exception as com_err:
            logger.error(f"COM status update failed: {com_err}")
        finally:
            if excel and should_quit:
                try:
                    excel.Quit()
                except Exception:
                    pass
                del excel


def skip_failed_loop_step(current_i: int, commands: List[Dict[str, Any]], loop_stack: List[Dict[str, Any]], workflow_id: str, update_ui: Optional[Callable[[str, bool], None]], err_msg: str) -> int:
    # Why: エラー発生時に行ステータスを更新し対応するloop_endまで安全にジャンプして次行へ継続
    curr_loop = loop_stack[-1]
    curr_loop["current_had_error"] = True
    logger.warning(f"[{workflow_id}] Step error: {err_msg}. continue_on_error enabled, skipping to next record.")
    if curr_loop.get("data_source") == "excel":
        curr_rec = curr_loop["records"][curr_loop["current_iteration"]]
        st_col = curr_loop.get("status_column")
        f_path = curr_loop.get("file_path")
        s_name = curr_loop.get("sheet_name")
        if st_col and f_path:
            write_excel_status(f_path, s_name, curr_rec.get("_row_idx"), st_col, f"エラー: {str(err_msg)[:25]}")
    if update_ui:
        update_ui(f"エラー行スキップ: {err_msg[:20]}", is_warning=True)
    nest = 1
    j = current_i + 1
    while j < len(commands) and nest > 0:
        if commands[j].get("method") == "loop_start":
            nest += 1
        elif commands[j].get("method") == "loop_end":
            nest -= 1
        if nest == 0:
            break
        j += 1
    return j


def get_excel_application(hwnd: Optional[int] = None, expected_title: str = "") -> Any:
    if platform.system() != "Windows":
        return None
    try:
        import win32gui
        import win32com.client
        import pythoncom

        excel7_hwnd = None
        if hwnd:
            # Why: タイトル文字列に依存せずウィンドウクラス名(XLMAIN/EXCEL7)でExcelを厳密判定
            cls_name = win32gui.GetClassName(hwnd)
            if cls_name == "EXCEL7":
                excel7_hwnd = hwnd
            elif cls_name == "XLMAIN":
                def enum_child(child: int, _: Any) -> bool:
                    nonlocal excel7_hwnd
                    try:
                        if win32gui.GetClassName(child) == "EXCEL7":
                            excel7_hwnd = child
                            return False
                    except Exception:
                        pass
                    return True
                try:
                    win32gui.EnumChildWindows(hwnd, enum_child, None)
                except Exception:
                    pass

        # 1. EXCEL7 から直接 COM オブジェクトを取得
        if excel7_hwnd:
            try:
                iid = (ctypes.c_byte * 16)(
                    0x00, 0x04, 0x02, 0x00, 0x00, 0x00, 0x00, 0x00,
                    0xC0, 0x00, 0x00, 0x00, 0x00, 0x00, 0x00, 0x46
                )
                p_acc = ctypes.c_void_p()
                hr = ctypes.oledll.oleacc.AccessibleObjectFromWindow(
                    excel7_hwnd, -16, iid, ctypes.byref(p_acc)
                )
                if hr == 0 and p_acc.value:
                    pycom_dll = ctypes.PyDLL(pythoncom.__file__)
                    pycom_func = pycom_dll.PyCom_PyObjectFromIUnknown
                    pycom_func.restype = ctypes.py_object
                    pycom_func.argtypes = (ctypes.c_void_p, ctypes.c_void_p, ctypes.c_int)
                    py_dispatch = pycom_func(p_acc.value, ctypes.byref(iid), 0)
                    if py_dispatch:
                        disp = win32com.client.Dispatch(py_dispatch)
                        app = getattr(disp, "Application", disp)
                        if app:
                            return app
            except Exception:
                pass

        # 2. フォールバック: 対象ブック名を特定して取得
        try:
            excel = win32com.client.GetActiveObject("Excel.Application")
            if excel:
                actual_title = ""
                if hwnd:
                    try:
                        actual_title = win32gui.GetWindowText(hwnd)
                    except Exception:
                        pass
                target_search = actual_title if actual_title else expected_title
                # Why: 操作対象ウィンドウの実タイトルから一致するWorkbookを特定して誤爆防止
                if target_search:
                    for wb in excel.Workbooks:
                        if wb.Name.lower() in target_search.lower():
                            wb.Activate()
                            return excel
                return excel
        except Exception:
            pass

    except Exception:
        pass
    return None


def execute_excel_action(args: Dict[str, Any], variables: Dict[str, Any], excel_app: Any, last_win_args: Optional[Dict[str, Any]], workflow_id: str) -> Tuple[Any, Any]:
    if platform.system() != "Windows":
        logger.warning(f"[{workflow_id}] Excel COM automation is only supported on Windows.")
        return excel_app, None

    import win32com.client
    import win32gui

    action = args.get("action", "")
    file_path = args.get("file_path")
    sheet_name = args.get("sheet_name")
    cell = args.get("cell")
    range_addr = args.get("range_address") or cell
    value = args.get("value")
    var_name = args.get("variable_name")
    macro_name = args.get("macro_name")
    close_save = args.get("close_save", True)

    if isinstance(value, str):
        for k, v in variables.items():
            value = value.replace(f"{{{{{k}}}}}", str(v)).replace(f"${{{k}}}", str(v))

    if excel_app is None:
        target_hwnd = (last_win_args.get("mapped_hwnd") if last_win_args else None) or win32gui.GetForegroundWindow()
        expected_t = last_win_args.get("window_title", "") if last_win_args else ""
        excel_app = get_excel_application(target_hwnd, expected_t)

    if excel_app is None and action == "open_workbook":
        excel_app = win32com.client.Dispatch("Excel.Application")
        excel_app.Visible = True

    if excel_app is None:
        try:
            excel_app = win32com.client.GetActiveObject("Excel.Application")
            excel_app.Visible = True
        except Exception:
            excel_app = win32com.client.Dispatch("Excel.Application")
            excel_app.Visible = True

    # Why: Excel操作時にウィンドウを前面化し他アプリとの切り替えを可視化
    try:
        if hasattr(excel_app, "Hwnd") and excel_app.Hwnd:
            ctypes.windll.user32.SetForegroundWindow(excel_app.Hwnd)
    except Exception:
        pass

    wb = None
    if file_path:
        resolved_path = str(Path(file_path).resolve())
        for open_wb in excel_app.Workbooks:
            if open_wb.FullName.lower() == resolved_path.lower() or open_wb.Name.lower() == Path(file_path).name.lower():
                wb = open_wb
                break
        if wb is None and action == "open_workbook":
            wb = excel_app.Workbooks.Open(resolved_path)

    if wb is None:
        try:
            wb = excel_app.ActiveWorkbook
        except Exception:
            wb = None

    sheet = None
    if wb:
        if sheet_name:
            try:
                sheet = wb.Sheets(sheet_name)
            except Exception:
                sheet = wb.ActiveSheet
        else:
            try:
                sheet = wb.ActiveSheet
            except Exception:
                pass

    result = None
    if action == "open_workbook":
        if wb:
            wb.Activate()
            result = wb.Name
    elif action == "save_workbook":
        if wb:
            if file_path and str(Path(file_path).resolve()).lower() != str(Path(wb.FullName).resolve()).lower():
                wb.SaveAs(str(Path(file_path).resolve()))
            else:
                wb.Save()
    elif action == "close_workbook":
        if wb:
            wb.Close(SaveChanges=close_save)
    elif action == "select_sheet":
        if sheet:
            sheet.Activate()
    elif action == "add_sheet":
        if wb:
            new_sheet = wb.Sheets.Add()
            if sheet_name:
                new_sheet.Name = sheet_name
            new_sheet.Activate()
    elif action == "read_cell":
        if sheet and cell:
            result = sheet.Range(cell).Value
            if var_name:
                variables[var_name] = result
                logger.info(f"[{workflow_id}] Read Excel cell {cell} -> variables['{var_name}'] = {result}")
    elif action == "write_cell":
        if sheet and cell:
            sheet.Range(cell).Value = value
    elif action == "read_range":
        if sheet and range_addr:
            raw_vals = sheet.Range(range_addr).Value
            result = [list(r) if isinstance(r, tuple) else r for r in raw_vals] if isinstance(raw_vals, tuple) else raw_vals
            if var_name:
                variables[var_name] = result
                logger.info(f"[{workflow_id}] Read Excel range {range_addr} -> variables['{var_name}']")
    elif action == "write_range":
        if sheet and range_addr:
            sheet.Range(range_addr).Value = value
    elif action == "insert_row":
        if sheet and cell:
            sheet.Rows(cell).Insert()
    elif action == "delete_row":
        if sheet and cell:
            sheet.Rows(cell).Delete()
    elif action == "clear_range":
        if sheet and range_addr:
            sheet.Range(range_addr).ClearContents()
    elif action == "run_macro":
        if macro_name:
            result = excel_app.Run(macro_name)
    elif action == "read_records":
        st_row = args.get("start_row", 2)
        end_r = args.get("end_row")
        st_col = args.get("status_column")
        skip_comp = args.get("skip_completed", True)
        records = read_excel_records(file_path, sheet_name, st_row, end_r, st_col, skip_comp)
        if var_name:
            variables[var_name] = records
        result = records
    elif action == "update_status":
        r_idx = args.get("row_index") or variables.get("_row_idx")
        st_col = args.get("status_column")
        st_val = str(args.get("value", "完了"))
        write_excel_status(file_path, sheet_name, r_idx, st_col, st_val)
        result = True
    else:
        logger.warning(f"[{workflow_id}] Unknown Excel action: {action}")

    return excel_app, result
