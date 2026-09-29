"""Module: @role: マクロエディタのキャンバス上でループ回数の設定、Excelデータ連携ダイアログ、および逆転警告を表示する補助UIウィジェット群。"""

from pathlib import Path
from PySide6.QtCore import Qt, Signal
from PySide6.QtGui import QColor, QPainter, QPainterPath, QPen
from PySide6.QtWidgets import (QFrame, QHBoxLayout, QLabel, QSpinBox, QWidget, 
                               QPushButton, QVBoxLayout, QFormLayout, 
                               QPushButton, QDialog, QVBoxLayout, QFormLayout, 
                               QLineEdit, QCheckBox, QComboBox, QFileDialog)


class WarningWidget(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setFixedSize(24, 24)
        self.setToolTip("ループの開始と終了が逆転しています。\n矢印線が逆転しないように注意してください。")

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)

        path = QPainterPath()
        path.moveTo(12, 2)
        path.lineTo(22, 20)
        path.lineTo(2, 20)
        path.closeSubpath()

        painter.setBrush(QColor("#ffcc00"))
        painter.setPen(QPen(QColor("#d13438"), 2))
        painter.drawPath(path)

        painter.setPen(Qt.PenStyle.NoPen)
        painter.setBrush(QColor("#d13438"))
        painter.drawRect(11, 8, 2, 6)
        painter.drawRect(11, 16, 2, 2)


class LoopSettingDialog(QDialog):
    def __init__(self, cmd: dict, workflow_dir: Path, parent=None):
        super().__init__(parent)
        self.cmd = cmd
        self.workflow_dir = workflow_dir
        self.args = cmd.setdefault("args", {})
        self.setWindowTitle("ループ設定")
        self.setMinimumWidth(440)
        self.setStyleSheet("""
            QDialog { background-color: #ffffff; }
            QLabel { color: #333333; font-size: 13px; }
            QLineEdit, QSpinBox, QComboBox { padding: 4px; border: 1px solid #d0d0d0; border-radius: 4px; font-size: 13px; }
        """)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(14)

        form = QFormLayout()

        self.mode_combo = QComboBox()
        self.mode_combo.addItem("固定回数ループ", "static")
        self.mode_combo.addItem("Excelデータ連携ループ", "excel")
        cur_mode = self.args.get("data_source", "static")
        self.mode_combo.setCurrentIndex(1 if cur_mode == "excel" else 0)
        form.addRow("ループ種別:", self.mode_combo)

        # 固定回数設定
        self.static_widget = QWidget()
        s_layout = QFormLayout(self.static_widget)
        s_layout.setContentsMargins(0, 0, 0, 0)
        self.count_spin = QSpinBox()
        self.count_spin.setRange(1, 99999)
        self.count_spin.setValue(self.args.get("loop_count", 10))
        s_layout.addRow("リピート回数:", self.count_spin)
        form.addRow(self.static_widget)

        # Excelデータ連携設定
        self.excel_widget = QWidget()
        e_layout = QFormLayout(self.excel_widget)
        e_layout.setContentsMargins(0, 0, 0, 0)

        file_box = QHBoxLayout()
        self.file_edit = QLineEdit(self.args.get("file_path", ""))
        self.file_edit.setPlaceholderText("Excelファイルパス...")
        file_box.addWidget(self.file_edit)

        browse_btn = QPushButton("参照...")
        browse_btn.setFixedWidth(60)
        browse_btn.clicked.connect(self._browse_file)
        file_box.addWidget(browse_btn)
        e_layout.addRow("Excelファイル:", file_box)

        self.sheet_edit = QLineEdit(self.args.get("sheet_name", ""))
        self.sheet_edit.setPlaceholderText("空欄でアクティブシート")
        e_layout.addRow("シート名:", self.sheet_edit)

        row_box = QHBoxLayout()
        self.start_row_spin = QSpinBox()
        self.start_row_spin.setRange(1, 99999)
        self.start_row_spin.setValue(self.args.get("start_row", 2))
        row_box.addWidget(QLabel("開始:"))
        row_box.addWidget(self.start_row_spin)

        self.end_row_spin = QSpinBox()
        self.end_row_spin.setRange(0, 99999)
        self.end_row_spin.setValue(self.args.get("end_row") or 0)
        self.end_row_spin.setSpecialValueText("末尾まで")
        row_box.addWidget(QLabel("終了:"))
        row_box.addWidget(self.end_row_spin)
        e_layout.addRow("対象行範囲:", row_box)

        self.status_col_edit = QLineEdit(self.args.get("status_column", ""))
        self.status_col_edit.setPlaceholderText("空欄で書込無 (例: F)")
        e_layout.addRow("ステータス記録列:", self.status_col_edit)

        self.skip_check = QCheckBox("完了ステータス済みの行をスキップ")
        self.skip_check.setChecked(bool(self.args.get("skip_completed", True)))
        e_layout.addRow("", self.skip_check)

        self.cont_check = QCheckBox("エラー発生時に中断せず次の行へ継続")
        self.cont_check.setChecked(bool(self.args.get("continue_on_error", False)))
        e_layout.addRow("", self.cont_check)

        form.addRow(self.excel_widget)

        def _toggle(idx):
            is_ex = idx == 1
            self.static_widget.setVisible(not is_ex)
            self.excel_widget.setVisible(is_ex)

        self.mode_combo.currentIndexChanged.connect(_toggle)
        _toggle(self.mode_combo.currentIndex())

        layout.addLayout(form)

        btn_box = QHBoxLayout()
        btn_box.addStretch()
        cancel_btn = QPushButton("キャンセル")
        cancel_btn.setStyleSheet("padding: 6px 14px; border: 1px solid #d0d0d0; border-radius: 4px;")
        cancel_btn.clicked.connect(self.reject)
        save_btn = QPushButton("決定")
        save_btn.setStyleSheet("background-color: #0078d4; color: white; font-weight: bold; padding: 6px 18px; border-radius: 4px;")
        save_btn.clicked.connect(self._save)
        btn_box.addWidget(cancel_btn)
        btn_box.addWidget(save_btn)
        layout.addLayout(btn_box)

    def _browse_file(self):
        fp, _ = QFileDialog.getOpenFileName(self, "Excelファイルを選択", str(self.workflow_dir), "Excel Files (*.xlsx *.xls *.xlsm);;All Files (*.*)")
        if fp:
            self.file_edit.setText(fp)

    def _save(self):
        is_ex = self.mode_combo.currentIndex() == 1
        self.args["data_source"] = "excel" if is_ex else "static"
        if is_ex:
            self.args["file_path"] = self.file_edit.text()
            self.args["sheet_name"] = self.sheet_edit.text()
            self.args["start_row"] = self.start_row_spin.value()
            ed = self.end_row_spin.value()
            self.args["end_row"] = ed if ed > 0 else None
            self.args["status_column"] = self.status_col_edit.text().strip().upper() or None
            self.args["skip_completed"] = self.skip_check.isChecked()
            self.args["continue_on_error"] = self.cont_check.isChecked()
        else:
            self.args["loop_count"] = self.count_spin.value()
        self.accept()


class LoopCountWidget(QFrame):
    count_changed = Signal(int, int)
    settings_changed = Signal()
    delete_requested = Signal(int)

    def __init__(self, loop_start_idx: int, cmd: dict, workflow_dir: Path, parent=None):
        super().__init__(parent)
        self.loop_start_idx = loop_start_idx
        self.cmd = cmd
        self.workflow_dir = workflow_dir
        self.args = cmd.setdefault("args", {})
        self.is_expanded = False

        self.setObjectName("LoopCountWidget")
        self._build_ui()
        self._update_appearance()

    def _build_ui(self):
        self.main_layout = QVBoxLayout(self)
        self.main_layout.setContentsMargins(10, 8, 10, 8)
        self.main_layout.setSpacing(8)

        # 1. サマリー表示バー
        self.bar_layout = QHBoxLayout()
        self.bar_layout.setContentsMargins(0, 0, 0, 0)
        self.bar_layout.setSpacing(6)

        self.icon_label = QLabel()
        self.bar_layout.addWidget(self.icon_label)

        self.summary_label = QLabel()
        self.summary_label.setStyleSheet("font-weight: bold; font-size: 12px; color: #333333;")
        self.bar_layout.addWidget(self.summary_label)

        # Why: 十分な幅と右パディングを確保し上下矢印が数値テキストに重なるのを完全防止
        self.spin_box = QSpinBox()
        self.spin_box.setRange(1, 99999)
        self.spin_box.setValue(self.args.get("loop_count", 10))
        self.spin_box.setFixedWidth(70)
        self.spin_box.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        self.spin_box.setStyleSheet("""
            QSpinBox {
                border: 1px solid #c7c7c7;
                border-radius: 4px;
                background-color: #ffffff;
                color: #0078d4;
                font-size: 12px;
                font-weight: bold;
                padding-right: 20px;
                padding-left: 4px;
            }
        """)
        self.spin_box.valueChanged.connect(self._on_spin_changed)
        self.bar_layout.addWidget(self.spin_box)

        self.unit_label = QLabel("回")
        self.unit_label.setStyleSheet("font-weight: bold; font-size: 12px; color: #333333;")
        self.bar_layout.addWidget(self.unit_label)

        self.bar_layout.addStretch()

        self.toggle_btn = QPushButton("詳細 ▼")
        self.toggle_btn.setFixedSize(56, 24)
        self.toggle_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.toggle_btn.setStyleSheet("""
            QPushButton {
                background: #f3f2f1;
                border: 1px solid #d0d0d0;
                border-radius: 3px;
                font-size: 11px;
                color: #444444;
            }
            QPushButton:hover {
                background: #e1dfdd;
                color: #0078d4;
            }
        """)
        self.toggle_btn.clicked.connect(self._toggle_expand)
        self.bar_layout.addWidget(self.toggle_btn)

        self.del_btn = QPushButton("✕")
        self.del_btn.setFixedSize(22, 22)
        self.del_btn.setToolTip("このループを解除する")
        self.del_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.del_btn.setStyleSheet("""
            QPushButton {
                background: transparent;
                border: none;
                color: #888888;
                font-size: 12px;
                font-weight: bold;
            }
            QPushButton:hover {
                color: #d13438;
            }
        """)
        self.del_btn.clicked.connect(lambda: self.delete_requested.emit(self.loop_start_idx))
        self.bar_layout.addWidget(self.del_btn)

        self.main_layout.addLayout(self.bar_layout)

        # 2. インライン詳細設定パネル（展開式）
        self.detail_panel = QWidget()
        self.detail_layout = QFormLayout(self.detail_panel)
        self.detail_layout.setContentsMargins(4, 8, 4, 4)
        self.detail_layout.setSpacing(8)
        self.detail_panel.setStyleSheet("""
            QLabel {
                font-size: 12px;
                color: #2b2b2b;
            }
            QLineEdit {
                font-size: 12px;
                padding: 4px 6px;
                border: 1px solid #c7c7c7;
                border-radius: 4px;
                background-color: #ffffff;
                color: #1f2937;
            }
            QLineEdit:focus {
                border: 1px solid #0078d4;
            }
            QSpinBox {
                font-size: 12px;
                padding: 3px 22px 3px 6px;
                border: 1px solid #c7c7c7;
                border-radius: 4px;
                background-color: #ffffff;
                color: #1f2937;
            }
            QSpinBox:focus {
                border: 1px solid #0078d4;
            }
            QComboBox {
                font-size: 12px;
                padding: 4px 8px;
                border: 1px solid #c7c7c7;
                border-radius: 4px;
                background-color: #ffffff;
                color: #1f2937;
            }
            QComboBox:hover {
                border: 1px solid #0078d4;
            }
            QComboBox QAbstractItemView {
                background-color: #ffffff;
                color: #1f2937;
                selection-background-color: #0078d4;
                selection-color: #ffffff;
                border: 1px solid #c7c7c7;
                outline: none;
            }
            QCheckBox {
                font-size: 12px;
                color: #2b2b2b;
            }
            QPushButton {
                font-size: 12px;
                color: #2b2b2b;
                background-color: #f3f2f1;
                border: 1px solid #c7c7c7;
                border-radius: 4px;
                padding: 3px 8px;
            }
            QPushButton:hover {
                background-color: #e1dfdd;
                color: #0078d4;
            }
        """)

        # モード選択
        self.mode_combo = QComboBox()
        self.mode_combo.addItem("固定回数ループ", "static")
        self.mode_combo.addItem("Excelデータ連携ループ", "excel")
        cur_mode = self.args.get("data_source", "static")
        self.mode_combo.setCurrentIndex(1 if cur_mode == "excel" else 0)
        self.mode_combo.currentIndexChanged.connect(self._on_mode_changed)
        self.detail_layout.addRow("種別:", self.mode_combo)

        # Excelグループ
        self.excel_group = QWidget()
        eg_layout = QFormLayout(self.excel_group)
        eg_layout.setContentsMargins(0, 0, 0, 0)
        eg_layout.setSpacing(7)

        file_box = QHBoxLayout()
        self.file_edit = QLineEdit(self.args.get("file_path", ""))
        self.file_edit.setPlaceholderText("Excelファイル...")
        self.file_edit.textChanged.connect(self._on_file_changed)
        file_box.addWidget(self.file_edit)

        browse_btn = QPushButton("参照")
        browse_btn.setFixedWidth(50)
        browse_btn.clicked.connect(self._browse_excel_file)
        file_box.addWidget(browse_btn)
        eg_layout.addRow("ファイル:", file_box)

        self.sheet_edit = QLineEdit(self.args.get("sheet_name", ""))
        self.sheet_edit.setPlaceholderText("空欄でアクティブシート")
        self.sheet_edit.textChanged.connect(self._on_sheet_changed)
        eg_layout.addRow("シート:", self.sheet_edit)

        # 検出列変数プレビュー
        self.cols_preview_label = QLabel()
        self.cols_preview_label.setWordWrap(True)
        self.cols_preview_label.setStyleSheet("font-size: 11px; color: #107c41; background: #eaf6ee; padding: 6px; border-radius: 4px;")
        eg_layout.addRow("検出列変数:", self.cols_preview_label)

        # Why: 1行詰め込み過密を解消し開始行と終了行を独立行にしてスピン矢印の被りを根本撲滅
        self.st_row_spin = QSpinBox()
        self.st_row_spin.setRange(1, 99999)
        self.st_row_spin.setValue(self.args.get("start_row", 2))
        self.st_row_spin.setFixedWidth(110)
        self.st_row_spin.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        self.st_row_spin.valueChanged.connect(lambda v: self._update_field("start_row", v))
        eg_layout.addRow("開始行:", self.st_row_spin)

        self.ed_row_spin = QSpinBox()
        self.ed_row_spin.setRange(0, 99999)
        self.ed_row_spin.setValue(self.args.get("end_row") or 0)
        self.ed_row_spin.setSpecialValueText("末尾まで")
        self.ed_row_spin.setFixedWidth(110)
        self.ed_row_spin.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        self.ed_row_spin.valueChanged.connect(lambda v: self._update_field("end_row", v if v > 0 else None))
        eg_layout.addRow("終了行:", self.ed_row_spin)

        self.status_col_edit = QLineEdit(self.args.get("status_column", ""))
        self.status_col_edit.setPlaceholderText("空欄で書込無 (例: F)")
        self.status_col_edit.setFixedWidth(110)
        self.status_col_edit.textChanged.connect(lambda v: self._update_field("status_column", v.upper()))
        eg_layout.addRow("ステータス列:", self.status_col_edit)

        self.skip_check = QCheckBox("完了行をスキップ")
        self.skip_check.setChecked(bool(self.args.get("skip_completed", True)))
        self.skip_check.toggled.connect(lambda v: self._update_field("skip_completed", v))
        eg_layout.addRow("", self.skip_check)

        self.cont_check = QCheckBox("エラー時次行継続")
        self.cont_check.setChecked(bool(self.args.get("continue_on_error", False)))
        self.cont_check.toggled.connect(lambda v: self._update_field("continue_on_error", v))
        eg_layout.addRow("", self.cont_check)

        self.detail_layout.addRow(self.excel_group)
        self.detail_panel.hide()
        self.main_layout.addWidget(self.detail_panel)

        self._refresh_columns_preview()

    def _refresh_columns_preview(self):
        f_path = self.args.get("file_path", "")
        s_name = self.args.get("sheet_name")
        cols_map = get_excel_columns_map(f_path, s_name, self.workflow_dir)
        self.args["columns_map"] = cols_map
        preview_texts = [f"{col}列:{name}" for col, name in list(cols_map.items())[:6]]
        if len(cols_map) > 6:
            preview_texts.append(f"...他{len(cols_map)-6}列")
        self.cols_preview_label.setText(" | ".join(preview_texts) if preview_texts else "列未検出")

    def _on_file_changed(self, v):
        self.args["file_path"] = v
        self._refresh_columns_preview()
        self._update_appearance()
        self.settings_changed.emit()

    def _on_sheet_changed(self, v):
        self.args["sheet_name"] = v
        self._refresh_columns_preview()
        self._update_appearance()
        self.settings_changed.emit()

    def _update_appearance(self):
        is_excel = self.args.get("data_source") == "excel"
        self.excel_group.setVisible(is_excel)

        if is_excel:
            self.icon_label.setText("📊")
            f_name = Path(self.args.get("file_path", "")).name if self.args.get("file_path") else "未選択"
            st_col = self.args.get("status_column") or "-"
            self.summary_label.setText(f"{f_name} [{st_col}列]")
            self.summary_label.setStyleSheet("font-weight: bold; font-size: 12px; color: #107c41;")
            self.spin_box.hide()
            self.unit_label.hide()
            self.setStyleSheet("""
                #LoopCountWidget {
                    background-color: #f6fcf8;
                    border: 2px solid #107c41;
                    border-radius: 8px;
                }
            """)
        else:
            self.icon_label.setText("🔁")
            self.summary_label.setText("固定ループ:")
            self.summary_label.setStyleSheet("font-weight: bold; font-size: 12px; color: #0078d4;")
            self.unit_label.setStyleSheet("font-weight: bold; font-size: 12px; color: #333333;")
            self.spin_box.show()
            self.unit_label.show()
            self.setStyleSheet("""
                #LoopCountWidget {
                    background-color: #f7faff;
                    border: 2px solid #0078d4;
                    border-radius: 8px;
                }
            """)
        self.adjustSize()

    def _toggle_expand(self):
        self.is_expanded = not self.is_expanded
        self.detail_panel.setVisible(self.is_expanded)
        self.toggle_btn.setText("閉じる ▲" if self.is_expanded else "詳細 ▼")
        # Why: 展開時幅を380pxに拡張し十分な余白と操作エリアを確保
        self.setFixedWidth(380 if self.is_expanded else 260)
        self.adjustSize()
        self.settings_changed.emit()

    def _on_mode_changed(self, idx):
        mode = "excel" if idx == 1 else "static"
        self.args["data_source"] = mode
        if mode == "excel":
            self._refresh_columns_preview()
        self._update_appearance()
        self.settings_changed.emit()

    def _on_spin_changed(self, val):
        self.args["loop_count"] = val
        self.count_changed.emit(self.loop_start_idx, val)

    def _update_field(self, key, val):
        self.args[key] = val
        self._update_appearance()
        self.settings_changed.emit()

    def _browse_excel_file(self):
        fp, _ = QFileDialog.getOpenFileName(
            self, "Excelファイルを選択", str(self.workflow_dir), "Excel Files (*.xlsx *.xls *.xlsm);;All Files (*.*)"
        )
        if fp:
            self.file_edit.setText(fp)
            self._on_file_changed(fp)
def get_excel_columns_map(file_path: str, sheet_name: str = None, base_dir: Path = None) -> dict[str, str]:
    """Excelの1行目を軽量スキャンし、列文字とヘッダー名の辞書を返す"""
    if not file_path:
        return {chr(65 + i): f"列{chr(65 + i)}" for i in range(8)}
    p = Path(file_path)
    if not p.exists() and base_dir and (base_dir / file_path).exists():
        p = base_dir / file_path
    if not p.exists():
        return {chr(65 + i): f"列{chr(65 + i)}" for i in range(8)}
    try:
        import openpyxl
        resolved = str(p.resolve())
        wb = openpyxl.load_workbook(resolved, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        cols = {}
        max_col = min(ws.max_column or 8, 26)
        for col_idx in range(1, max_col + 1):
            letter = openpyxl.utils.get_column_letter(col_idx)
            val = ws.cell(row=1, column=col_idx).value
            cols[letter] = str(val).strip() if val is not None else f"列{letter}"
        wb.close()
        return cols if cols else {chr(65 + i): f"列{chr(65 + i)}" for i in range(8)}
    except Exception:
        return {chr(65 + i): f"列{chr(65 + i)}" for i in range(8)}
